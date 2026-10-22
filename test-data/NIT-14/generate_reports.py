#!/usr/bin/env python3
"""
Generate all 9 PWD tender report .docx files for NIT-14/2026-27.

Input:  test-data/NIT-14/Inputs/openingnit14.pdf (already parsed)
Output: test-data/NIT-14/outputs/*.docx
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# ── NIT-14 Data (extracted from openingnit14.pdf) ──────────────────────

NIT_NO = "14/2026-27"
NIT_DATE = "21-Oct-2026"
OFFICE = "OFFICE OF THE EXECUTIVE ENGINEER, P.W.D. DISTRICT DIVISION-II, UDAIPUR"
OFFICE_SHORT = "Executive Engineer, P.W.D. District Division-II, Udaipur"
SIGN_AUTH_L1 = "[ ANIL KHINCHI ]"
SIGN_AUTH_L2 = "Executive Engineer"
SIGN_AUTH_L3 = "PWD District Division-II, Udaipur"

WORKS = [
    {
        "sno": 1,
        "name": "VARIOUS REPAIR AND MAINTENANCE WORKS AT GOVT. TRIBAL MODEL PUBLIC RESIDENTIAL SCHOOL DHINKLI (BADGAON) UNDER DHARTI ABHA JANJATIYA GRAM UTKRASH ABHIYAN 2026-27",
        "ubn": "",
        "g_schedule": 6475245.91,
        "bidder_name": "M/s UDAIPUR BUILDCOM",
        "bid_amount": 5426903.60,
        "bid_rate_pct": 16.19,
        "bid_rate_type": "below",
        "period": "",
        "status": "accepted",
        "tenders_sold": 3,
        "tenders_received": 3,
        "responsive_bidders": 3,
        "bsr": "PWD Roads BSR 2022",
        "single_tender": False,
    },
    {
        "sno": 2,
        "name": "REPAIR AND MAINTENANCE WORK OF GOVT. GIRLS ASHRAM HOSTEL AT PIPLI KHEDA (GOGUNDA) UNDER DHARTI ABHA JANJATIYA GRAM UTKRASH ABHIYAN 2026-27",
        "ubn": "",
        "g_schedule": 5295360.43,
        "bidder_name": "M/s UDAIPUR BUILDCOM",
        "bid_amount": 4496290.54,
        "bid_rate_pct": 15.09,
        "bid_rate_type": "below",
        "period": "",
        "status": "accepted",
        "tenders_sold": 4,
        "tenders_received": 4,
        "responsive_bidders": 4,
        "bsr": "PWD Roads BSR 2022",
        "single_tender": False,
    },
    {
        "sno": 3,
        "name": "WATER FACILITY WORKS OF GOVT. TRIBAL BOYS ASHRAM HOSTEL AT RAWACHH (SAYRA) UNDER DHARTI ABHA JANJATIYA GRAM UTKRASH ABHIYAN 2026-27",
        "ubn": "",
        "g_schedule": 431337.81,
        "bidder_name": "M/s PALIWAL INFRA PROJECTS",
        "bid_amount": 418182.01,
        "bid_rate_pct": 3.05,
        "bid_rate_type": "below",
        "period": "",
        "status": "accepted",
        "tenders_sold": 1,
        "tenders_received": 1,
        "responsive_bidders": 1,
        "bsr": "PWD Roads BSR 2022",
        "single_tender": True,
    },
]

# ── Computation helpers ───────────────────────────────────────────────

ONES = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine",
        "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen",
        "Seventeen", "Eighteen", "Nineteen"]
TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]

def convert_hundreds(n):
    if n == 0: return ""
    if n < 20: return ONES[n]
    if n < 100:
        return TENS[n // 10] + (" " + ONES[n % 10] if n % 10 else "")
    return ONES[n // 100] + " Hundred" + (" " + convert_hundreds(n % 100) if n % 100 else "")

def number_to_indian_words(amount):
    if amount == 0: return "Zero Only"
    amt = round(amount)
    crore = amt // 10000000
    lakh = (amt % 10000000) // 100000
    thousand = (amt % 100000) // 1000
    rest = amt % 1000
    words = ""
    if crore: words += convert_hundreds(crore) + " Crore "
    if lakh: words += convert_hundreds(lakh) + " Lakh "
    if thousand: words += convert_hundreds(thousand) + " Thousand "
    if rest: words += convert_hundreds(rest)
    return words.strip() + " Only"

def compute_stamp_duty(bid_amount):
    if bid_amount <= 5000000: return 1000
    return min(round(bid_amount * 0.0015), 2500000)

def compute_aps(g_schedule, bid_amount, bid_rate_type):
    if bid_rate_type != "below": return 0
    if not g_schedule or g_schedule <= 0: return 0
    pct_below = ((g_schedule - bid_amount) / g_schedule) * 100
    if pct_below < 15: return 0
    excess = pct_below - 15
    return round(0.5 * (excess / 100) * g_schedule)

def get_authority(bid_amount):
    if bid_amount <= 7500000:
        return "Executive Engineer, P.W.D. District Division-II, Udaipur"
    elif bid_amount <= 30000000:
        return "Superintending Engineer, PWD City Circle, Udaipur"
    elif bid_amount <= 75000000:
        return "Additional Chief Engineer, PWD Zone Udaipur"
    elif bid_amount <= 150000000:
        return "Chief Engineer, PWD Rajasthan"
    else:
        return "EB/TAC — Full Powers"

def format_inr(amount):
    if amount is None: return "Nil"
    return f"Rs. {amount:,.2f}"

def format_inr_lacs(amount):
    if amount is None: return "Nil"
    lacs = amount / 100000
    return f"Rs. {lacs:,.2f} Lacs"

# ── Document helpers ──────────────────────────────────────────────────

FONT = "Times New Roman"
FS_BODY = Pt(12)
FS_SMALL = Pt(10)
FS_HEADER = Pt(15)

def setup_a4(doc):
    """Set A4 page size with PWD-standard margins (matching doc-generator.ts).
    Also explicitly suppress any footer so nothing prints below the content
    when the file is opened in Word, Google Docs, or any browser viewer.
    """
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.27)      # 0.5 inch  (720 twips)
    section.bottom_margin = Cm(1.27)  # 0.5 inch  (720 twips)
    section.left_margin = Cm(1.905)   # 0.75 inch (1080 twips)
    section.right_margin = Cm(1.905) # 0.75 inch (1080 twips)

    # ── Explicitly suppress footer ──────────────────────────────────
    # Unlink from any previous section and leave the footer empty so
    # no footer text, page number, URL, or system name ever prints.
    footer = section.footer
    footer.is_linked_to_previous = False
    # Clear any existing footer paragraphs (python-docx may inherit defaults)
    for p in footer.paragraphs:
        p.clear()

    # ── Set neutral document core properties ────────────────────────
    # Some viewers (especially browsers) print the document title as a
    # header. Set it to empty so nothing leaks.
    core = doc.core_properties
    core.title = ""
    core.author = ""
    core.comments = ""
    core.category = ""

def set_cell_text(cell, text, bold=False, size=FS_BODY):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = size
    run.bold = bold

def add_header(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = FS_HEADER
    run.bold = True

def add_centered(doc, text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = FS_BODY
    run.bold = bold

def add_para(doc, text, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(0.5)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = FS_BODY
    run.bold = bold

def add_justified(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = FS_BODY

def add_sign(doc, l1, l2, l3):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(l1)
    run.font.name = FONT
    run.bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(l2)
    run.font.name = FONT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(l3)
    run.font.name = FONT

def add_empty(doc):
    doc.add_paragraph()

def add_page_break(doc):
    doc.add_page_break()

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        elem = OxmlElement(f'w:{edge}')
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), '4')
        elem.set(qn('w:color'), '000000')
        borders.append(elem)
    tblPr.append(borders)

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "outputs")
os.makedirs(OUTPUT_DIR, exist_ok=True)
NIT_SAFE = NIT_NO.replace("/", "-")

# ── 1. Scrutiny Note Sheet ────────────────────────────────────────────

def gen_scrutiny():
    doc = Document()
    setup_a4(doc)
    for work in WORKS:
        computed_stamp = compute_stamp_duty(work["bid_amount"])
        computed_aps = compute_aps(work["g_schedule"], work["bid_amount"], work["bid_rate_type"])
        authority = get_authority(work["bid_amount"])
        rate_str = f'{work["bid_rate_pct"]:.2f}% Below {work["bsr"]}'

        add_header(doc, OFFICE)
        add_empty(doc)
        add_centered(doc, "TENDER SCRUTINY / NOTE SHEET", bold=True)
        add_centered(doc, f"NIT No. {NIT_NO}   |   S.No. {work['sno']}", bold=True)
        add_empty(doc)
        add_para(doc, f"NAME OF WORK: {work['name']}", bold=True)
        add_empty(doc)

        rows = [
            ("1", "Administration and Financial Sanction (Ref. and Amt.)", ""),
            ("2", "Technical Sanction (Ref. & Amt.)", ""),
            ("3", "Ref. & Amount Deposited in Case of Deposit Work", "Not Applicable"),
            ("4", "NIT Issued (No. & Date)", f"NIT No. {NIT_NO} Dated {NIT_DATE}"),
            ("5", "DIPR Ref. for Publication", ""),
            ("6", "UBN No.", work["ubn"] or "—"),
            ("7", "Whether it is a Short-Term NIT (Y/N)", "No"),
            ("8", "Whether Permission of Short NIT Obtained", "Not Applicable"),
            ("9", "NIT Amount of Work", format_inr_lacs(work["g_schedule"])),
            ("10", "Last Date of Online Submission and Time", ""),
            ("11", "Date of Online Opening of Tender", NIT_DATE),
            ("12", "Ref. of Corrigendum (if any)", "Not Applicable"),
            ("13", "Revised Date of Submission and Time (if any)", "Not Applicable"),
            ("14", "Revised Date of Opening (if any)", "Not Applicable"),
            ("15", "No. of Tenders Sold", str(work["tenders_sold"])),
            ("16", "No. of Tenders Received", str(work["tenders_received"])),
            ("17", "Date of Technical Bid Opening", "Not Applicable"),
            ("18", "No. of Responsive Bidders", str(work["responsive_bidders"])),
            ("19", "Date of Financial Bid Opening", NIT_DATE),
            ("20", "BSR", work["bsr"]),
            ("21", "G Schedule Amount", format_inr(work["g_schedule"])),
            ("22", "Name of Lowest Bidder", work["bidder_name"]),
            ("23", "Class of Contractor", ""),
            ("24", "Tender Premium Quoted by the Lowest Bidder", rate_str),
            ("25", "Bid Amount of the Lowest Bidder", format_inr(work["bid_amount"])),
            ("26", "Lowest Rate with Condition", f'{work["bid_rate_pct"]:.2f}% Below, No Condition'),
            ("27", "Financial Implication of Condition", "Not Applicable"),
            ("28", "EMD & Other Deposit Details", "Enclosed"),
            ("29", "Validity of Tender", "90 Days"),
            ("30", "Authority Competent to Accept", authority),
            ("31", "Time of Completion", work["period"] or "—"),
            ("32", "NIT Publication Details", "Enclosed"),
            ("33", "Recommendation", "Enclosed"),
        ]

        table = doc.add_table(rows=len(rows), cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        for i, (sno, desc, details) in enumerate(rows):
            set_cell_text(table.cell(i, 0), sno, size=FS_SMALL)
            set_cell_text(table.cell(i, 1), desc, size=FS_SMALL)
            set_cell_text(table.cell(i, 2), details, size=FS_SMALL)

        add_empty(doc)
        add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)
        if work["sno"] < len(WORKS):
            add_page_break(doc)

    path = os.path.join(OUTPUT_DIR, f"01-Scrutiny-Sheet-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── 2. Acceptance Letters ─────────────────────────────────────────────

def gen_acceptance():
    doc = Document()
    setup_a4(doc)
    for work in WORKS:
        stamp_duty = compute_stamp_duty(work["bid_amount"])
        aps = compute_aps(work["g_schedule"], work["bid_amount"], work["bid_rate_type"])
        bid_words = number_to_indian_words(work["bid_amount"])
        rate_str = f'{work["bid_rate_pct"]:.2f}% Below the schedule \'G\' amount'

        add_header(doc, OFFICE)
        add_empty(doc)
        p = doc.add_paragraph()
        run = p.add_run("No.:                              \tDate:                              ")
        run.font.name = FONT; run.font.size = FS_BODY
        add_empty(doc)
        add_centered(doc, "(Letter of Acceptance of Tender)", bold=True)
        add_empty(doc)
        add_para(doc, f'{work["bidder_name"]},')
        add_para(doc, "___________________________________.")
        add_empty(doc)
        add_para(doc, f"  Sub :- {work['name']}", bold=False)
        add_para(doc, f"  Ref :- NIT No. {NIT_NO}", bold=False)
        add_empty(doc)
        add_para(doc, "Sir,")
        add_justified(doc, f'Your Tender for the above work has been accepted on behalf of the Governor of Rajasthan as per rule @{rate_str} to Rs. {work["bid_amount"]:,.2f}/- ({bid_words}) by the {get_authority(work["bid_amount"])} Vide NIT No. {NIT_NO} Dated {NIT_DATE}.')
        add_justified(doc, f'Security Deposit as per rule of the gross amount of the running bill shall be deducted from each running bill or you may opt to deposit full amount of security deposit in the share of bank guarantee or any acceptable form of security before or at the time of executing agreement. Kindly submit the required stamp duty of Rs. {stamp_duty:,}/- as per rule and Deposit Additional Performance Security Amounting to Rs. {aps:,}.00 in this Office and do Sign the agreement within 10 days failing which action as per rule may be initiated.')
        add_justified(doc, "The receipt of the may please be acknowledged.")
        add_empty(doc)
        add_para(doc, f"Stamp Duty = Rs. {stamp_duty:,}/-", bold=True)
        add_para(doc, f"APS Amount  = Rs. {aps:,}/-" if aps else "APS Amount  = NIL", bold=True)
        add_empty(doc)
        add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)
        add_empty(doc)
        p = doc.add_paragraph()
        run = p.add_run("No.:                              \tDate:                              ")
        run.font.name = FONT; run.font.size = FS_BODY
        add_para(doc, "Copy submitted to the following :-")
        add_para(doc, "1-  The Additional Chief Engineer PWD Zone Udaipur.")
        add_para(doc, "2-  The Superintending Engineer PWD City Circle Udaipur.")
        add_para(doc, "3-  Assistant Engineer, P.W.D. Sub Dn.")
        add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)
        if work["sno"] < len(WORKS):
            add_page_break(doc)

    path = os.path.join(OUTPUT_DIR, f"02-Acceptance-Letters-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── 3. Work Orders ───────────────────────────────────────────────────

def gen_work_orders():
    doc = Document()
    setup_a4(doc)
    for work in WORKS:
        stamp_duty = compute_stamp_duty(work["bid_amount"])
        bid_words = number_to_indian_words(work["bid_amount"])
        rate_str = f'{work["bid_rate_pct"]:.2f}% Below B.S.R. 2022 (Roads)'

        add_header(doc, "OFFICE OF THE EXECUTIVE ENGINEER P.W.D. DISTT. DN. II UDAIPUR")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("No :-                              Date :-                    ")
        run.font.name = FONT; run.font.size = FS_BODY
        add_empty(doc)
        add_centered(doc, "(WRITTEN ORDER TO COMMENCEMENT OF WORK)", bold=True)
        add_empty(doc)
        add_para(doc, f'M/s {work["bidder_name"]},')
        add_para(doc, "___________________________________")
        add_empty(doc)
        add_para(doc, f"  Sub :- {work['name']}")
        add_para(doc, f"  Ref :- NIT No. {NIT_NO}")
        add_empty(doc)
        add_para(doc, "Sir,")
        add_justified(doc, f'With reference to the above noted subject and reference, you are hereby directed to start the work immediately and complete the same within the stipulated period. The work has been awarded to you at {rate_str} amounting to Rs. {work["bid_amount"]:,.2f}/- ({bid_words}).')
        add_justified(doc, f'You are also directed to submit the stamp duty of Rs. {stamp_duty:,}/- and sign the agreement within 10 days from the date of issue of this order.')
        add_justified(doc, "The receipt of this order may please be acknowledged.")
        add_empty(doc)
        add_para(doc, "Copy submitted to the following :-")
        add_para(doc, "1-  The Additional Chief Engineer PWD Zone Udaipur.")
        add_para(doc, "2-  The Superintending Engineer PWD City Circle Udaipur.")
        add_para(doc, "3-  Assistant Engineer, P.W.D. Sub Dn.")
        add_empty(doc)
        add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)
        if work["sno"] < len(WORKS):
            add_page_break(doc)

    path = os.path.join(OUTPUT_DIR, f"03-Work-Orders-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── 4. Negotiation Letters ───────────────────────────────────────────

def gen_negotiation_letters():
    doc = Document()
    setup_a4(doc)
    for work in WORKS:
        add_header(doc, OFFICE)
        add_empty(doc)
        p = doc.add_paragraph()
        run = p.add_run("No.: _______________________                    Date: _______________________")
        run.font.name = FONT; run.font.size = FS_BODY
        add_empty(doc)
        add_centered(doc, "(NEGOTIATIONS)", bold=True)
        add_empty(doc)
        add_para(doc, f'{work["bidder_name"]},')
        add_para(doc, "Udaipur.")
        add_empty(doc)
        add_para(doc, f'  Sub :- {work["name"]}')
        add_para(doc, f"  Ref :- NIT No. {NIT_NO}")
        add_empty(doc)
        add_para(doc, "Sir,")
        add_justified(doc, f'Please refer to your tender submitted for the above work, since the rates offered by you have been considered on higher side. It has been decided to hold the negotiation on _____________.')
        add_justified(doc, "You are therefore requested to please quote your negotiated offer in a sealed cover which shall be opened on __________ at __ P.M. in this office in presence of the contractors who choose to be present on date.")
        add_justified(doc, "It may please be carefully noted that the condition what so ever if laid may invalid at your offer.")
        add_empty(doc)
        add_justified(doc, "The receipt of the may please be acknowledged.")
        add_empty(doc)
        add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)
        if work["sno"] < len(WORKS):
            add_page_break(doc)

    path = os.path.join(OUTPUT_DIR, f"04-Negotiation-Letters-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── 5. Bank BG/FDR Letters ───────────────────────────────────────────

def gen_bank_bg_letters():
    doc = Document()
    setup_a4(doc)
    for work in WORKS:
        stamp_duty = compute_stamp_duty(work["bid_amount"])
        aps = compute_aps(work["g_schedule"], work["bid_amount"], work["bid_rate_type"])
        security = round(work["bid_amount"] * 0.05)  # 5% security deposit estimate
        total_deposit = stamp_duty + aps + security

        add_header(doc, OFFICE)
        add_empty(doc)
        p = doc.add_paragraph()
        run = p.add_run("No.: _______________________                    Date: _______________________")
        run.font.name = FONT; run.font.size = FS_BODY
        add_empty(doc)
        add_centered(doc, "To, The Branch Manager,", bold=True)
        add_para(doc, "___________________________________,")
        add_para(doc, "Udaipur.")
        add_empty(doc)
        add_para(doc, f"  Sub :- Confirmation of Bank Guarantee / FDR submitted by {work['bidder_name']}")
        add_para(doc, f"  Ref :- NIT No. {NIT_NO}, S.No. {work['sno']}")
        add_empty(doc)
        add_para(doc, "Sir,")
        add_justified(doc, f'With reference to the above noted subject, it is to inform you that {work["bidder_name"]} has participated in the tender for the work "{work["name"]}" and has been declared as the lowest bidder. The said contractor has submitted a Bank Guarantee / FDR No. _______________ dated _______________ for Rs. {security:,}/- as Security Deposit.')
        add_justified(doc, "You are requested to kindly confirm the genuineness and validity of the above mentioned Bank Guarantee / FDR at the earliest so that further necessary action may be taken at this end.")
        add_justified(doc, "The receipt of the may please be acknowledged.")
        add_empty(doc)
        add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)
        if work["sno"] < len(WORKS):
            add_page_break(doc)

    path = os.path.join(OUTPUT_DIR, f"05-Bank-BG-FDR-Letters-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── 6. Negotiation Reply Format ──────────────────────────────────────

def gen_negotiation_reply():
    doc = Document()
    setup_a4(doc)
    for work in WORKS:
        add_para(doc, "To,")
        add_para(doc, "EXECUTIVE ENGINEER,")
        add_para(doc, "P.W.D. DISTRICT DIVISION-II,")
        add_para(doc, "UDAIPUR")
        add_empty(doc)
        add_para(doc, f"Subject: {work['name']}")
        add_empty(doc)
        add_para(doc, f"Reference: NIT No. {NIT_NO}")
        add_empty(doc)
        add_para(doc, "Dear Sir,")
        add_justified(doc, "With reference to above subject, we would like to inform you that we have submitted our price on very competitive rates. Even then we would like to offer some extra discount as under:")
        add_empty(doc)
        add_justified(doc, "1   Our negotiated rate shall be __% Below the G-Schedule amount (In words Rupees ________________________ only).")
        add_empty(doc)
        add_justified(doc, "Hope you will find the above in order & arrange to release the work order at your earliest possible & oblige.")
        add_empty(doc)
        add_justified(doc, "Thanking you and assuring you of our best services always.")
        add_empty(doc)
        add_para(doc, "Yours faithfully,")
        add_empty(doc)
        add_empty(doc)
        add_para(doc, f'{work["bidder_name"]},')
        add_para(doc, "Udaipur.")
        add_para(doc, "Mobile: _________________")
        if work["sno"] < len(WORKS):
            add_page_break(doc)

    path = os.path.join(OUTPUT_DIR, f"06-Negotiation-Reply-Format-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── 7. Challan Verification Sheet ─────────────────────────────────────

def gen_challan_verification():
    doc = Document()
    setup_a4(doc)
    add_header(doc, OFFICE)
    add_empty(doc)
    add_centered(doc, "eGROSS CHALLAN VERIFICATION SHEET", bold=True)
    add_centered(doc, f"NIT No. {NIT_NO}", bold=True)
    add_empty(doc)
    add_para(doc, f"NIT Date: {NIT_DATE}")
    add_empty(doc)

    headers = ["S.No", "Work S.No", "GRN No.", "Challan No.", "Vendor Name", "Head 0075 Amount", "Head 8443 Amount"]
    data = [
        ["1", "1", "—", "—", "—", "—", "—"],
        ["2", "2", "—", "—", "—", "—", "—"],
        ["3", "3", "—", "—", "—", "—", "—"],
    ]

    table = doc.add_table(rows=1 + len(data), cols=len(headers))
    set_table_borders(table)
    for i, h in enumerate(headers):
        set_cell_text(table.cell(0, i), h, bold=True, size=FS_SMALL)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            set_cell_text(table.cell(r + 1, c), val, size=FS_SMALL)

    add_empty(doc)
    add_para(doc, "Note: EMD and fee challan details to be filled from eGross portal / EMD deposit document.")
    add_empty(doc)
    add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)

    path = os.path.join(OUTPUT_DIR, f"07-Challan-Verification-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── 8. Publication Cost Statement ─────────────────────────────────────

def gen_publication_cost():
    doc = Document()
    setup_a4(doc)
    add_header(doc, OFFICE)
    add_empty(doc)
    add_centered(doc, "NIT PUBLICATION COST STATEMENT", bold=True)
    add_centered(doc, f"NIT No. {NIT_NO}", bold=True)
    add_empty(doc)
    add_para(doc, f"NIT Date: {NIT_DATE}")
    add_empty(doc)

    headers = ["S.No", "Newspaper Name", "Edition", "Advt. Size", "Sq. Cm.", "Rate", "Amount"]
    data = [["1", "—", "—", "—", "—", "—", "—"]]

    table = doc.add_table(rows=1 + len(data), cols=len(headers))
    set_table_borders(table)
    for i, h in enumerate(headers):
        set_cell_text(table.cell(0, i), h, bold=True, size=FS_SMALL)
    for r, row in enumerate(data):
        for c, val in enumerate(row):
            set_cell_text(table.cell(r + 1, c), val, size=FS_SMALL)

    add_empty(doc)
    add_para(doc, "DIPR RO No.: _________________")
    add_para(doc, "Release Date: _________________")
    add_para(doc, "Total Publication Cost: Rs. _________________")
    add_empty(doc)
    add_para(doc, "Note: Publication details to be filled from DIPR publication document.")
    add_empty(doc)
    add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)

    path = os.path.join(OUTPUT_DIR, f"08-Publication-Cost-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── 9. Master Complete Record ────────────────────────────────────────

def gen_master_record():
    doc = Document()
    setup_a4(doc)
    add_header(doc, OFFICE)
    add_empty(doc)
    add_centered(doc, "MASTER COMPLETE RECORD", bold=True)
    add_centered(doc, f"NIT No. {NIT_NO} Dated {NIT_DATE}", bold=True)
    add_empty(doc)

    headers = ["S.No", "Name of Work", "UBN", "G-Schedule Amount", "Lowest Bidder", "Bid Amount", "Rate %", "Stamp Duty", "APS", "Status"]
    table = doc.add_table(rows=1 + len(WORKS), cols=len(headers))
    set_table_borders(table)
    for i, h in enumerate(headers):
        set_cell_text(table.cell(0, i), h, bold=True, size=FS_SMALL)

    for r, work in enumerate(WORKS):
        stamp = compute_stamp_duty(work["bid_amount"])
        aps = compute_aps(work["g_schedule"], work["bid_amount"], work["bid_rate_type"])
        row_data = [
            str(work["sno"]),
            work["name"][:60] + "..." if len(work["name"]) > 60 else work["name"],
            work["ubn"] or "—",
            format_inr(work["g_schedule"]),
            work["bidder_name"],
            format_inr(work["bid_amount"]),
            f'{work["bid_rate_pct"]:.2f}% Below',
            f"Rs. {stamp:,}/-",
            f"Rs. {aps:,}/-" if aps else "NIL",
            work["status"],
        ]
        for c, val in enumerate(row_data):
            set_cell_text(table.cell(r + 1, c), val, size=FS_SMALL)

    add_empty(doc)

    # Totals
    total_bid = sum(w["bid_amount"] for w in WORKS)
    total_stamp = sum(compute_stamp_duty(w["bid_amount"]) for w in WORKS)
    total_aps = sum(compute_aps(w["g_schedule"], w["bid_amount"], w["bid_rate_type"]) for w in WORKS)
    add_para(doc, f"Total Bid Amount: Rs. {total_bid:,.2f}/-", bold=True)
    add_para(doc, f"Total Stamp Duty: Rs. {total_stamp:,}/-", bold=True)
    add_para(doc, f"Total APS: Rs. {total_aps:,}/-", bold=True)
    add_empty(doc)
    add_sign(doc, SIGN_AUTH_L1, SIGN_AUTH_L2, SIGN_AUTH_L3)

    path = os.path.join(OUTPUT_DIR, f"09-Master-Record-NIT-{NIT_SAFE}.docx")
    doc.save(path)
    print(f"  Generated: {os.path.basename(path)}")

# ── Main ──────────────────────────────────────────────────────────────

if __name__ == "__main__":
    print(f"Generating NIT-14/2026-27 reports to: {OUTPUT_DIR}")
    print(f"  Works: {len(WORKS)}")
    for w in WORKS:
        stamp = compute_stamp_duty(w["bid_amount"])
        aps = compute_aps(w["g_schedule"], w["bid_amount"], w["bid_rate_type"])
        auth = get_authority(w["bid_amount"])
        print(f"  S.No.{w['sno']}: G={w['g_schedule']:,.2f}  Bid={w['bid_amount']:,.2f}  "
              f"StampDuty={stamp:,}  APS={aps:,}  Auth={auth.split(',')[0]}")
    print()

    gen_scrutiny()
    gen_acceptance()
    gen_work_orders()
    gen_negotiation_letters()
    gen_bank_bg_letters()
    gen_negotiation_reply()
    gen_challan_verification()
    gen_publication_cost()
    gen_master_record()

    print(f"\nAll 9 reports generated successfully in {OUTPUT_DIR}")
